"""Re-import Partivia quotes from MS365 PARTIVIA SPAGNA folder.

Downloads all emails, re-parses them with the updated Claude prompt (English),
and saves each email linked to its quotes in the DB.

Requires:
  - Flask app running at http://localhost:5000
  - MS365 Graph API credentials configured in saba-form
"""

import asyncio
import base64
import logging
import sys
from dataclasses import dataclass, field
from pathlib import Path

import httpx

# Add saba-form/mail_digest/src to path for GraphClient
MAIL_DIGEST = Path(__file__).parent.parent / "saba-form" / "mail_digest"
sys.path.insert(0, str(MAIL_DIGEST / "src"))

from mail_digest.graph.client import GraphClient  # noqa: E402

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)

MAILBOX = "info@sabae20.it"
TARGET_FOLDER = "PARTIVIA SPAGNA"
API_BASE = "http://localhost:5000"


# ── Data structures ──────────────────────────────────────────────


@dataclass
class FolderInfo:
    id: str
    name: str
    parent_path: str
    message_count: int = 0

    @property
    def full_path(self) -> str:
        return f"{self.parent_path} / {self.name}" if self.parent_path else self.name


@dataclass
class MessageData:
    id: str
    subject: str
    sender_name: str
    sender_email: str
    date: str
    body_text: str
    folder_path: str
    has_attachments: bool
    attachments: list[dict] = field(default_factory=list)


# ── Graph API helpers (from partivia_quotes.py) ──────────────────


async def find_partivia_folder(client: GraphClient) -> FolderInfo | None:
    data = await client.get(
        f"/users/{MAILBOX}/mailFolders",
        params={"$top": "100"},
    )
    for f in data.get("value", []):
        if TARGET_FOLDER.lower() in f["displayName"].lower():
            return FolderInfo(
                id=f["id"], name=f["displayName"], parent_path="",
                message_count=f.get("totalItemCount", 0),
            )
        children = await client.get(
            f"/users/{MAILBOX}/mailFolders/{f['id']}/childFolders",
            params={"$top": "100"},
        )
        for c in children.get("value", []):
            if TARGET_FOLDER.lower() in c["displayName"].lower():
                return FolderInfo(
                    id=c["id"], name=c["displayName"], parent_path=f["displayName"],
                    message_count=c.get("totalItemCount", 0),
                )
    return None


async def scan_subfolders(client: GraphClient, folder: FolderInfo) -> list[FolderInfo]:
    all_folders = []
    if folder.message_count > 0:
        all_folders.append(folder)
    data = await client.get(
        f"/users/{MAILBOX}/mailFolders/{folder.id}/childFolders",
        params={"$top": "100"},
    )
    for child in data.get("value", []):
        child_info = FolderInfo(
            id=child["id"], name=child["displayName"],
            parent_path=folder.full_path,
            message_count=child.get("totalItemCount", 0),
        )
        child_results = await scan_subfolders(client, child_info)
        all_folders.extend(child_results)
    return all_folders


async def fetch_messages(client: GraphClient, folder: FolderInfo) -> list[MessageData]:
    from bs4 import BeautifulSoup

    messages = []
    url = f"/users/{MAILBOX}/mailFolders/{folder.id}/messages"
    params = {
        "$select": "id,subject,from,receivedDateTime,body,hasAttachments",
        "$orderby": "receivedDateTime desc",
        "$top": "50",
    }
    while url:
        if url.startswith("https://"):
            data = await client.get_absolute(url)
        else:
            data = await client.get(url, params=params)
        params = None

        for msg in data.get("value", []):
            sender = msg.get("from", {}).get("emailAddress", {})
            body_html = msg.get("body", {}).get("content", "")
            soup = BeautifulSoup(body_html, "html.parser")
            body_text = soup.get_text(separator="\n", strip=True)

            messages.append(MessageData(
                id=msg["id"],
                subject=msg.get("subject", "(no subject)"),
                sender_name=sender.get("name", ""),
                sender_email=sender.get("address", ""),
                date=msg.get("receivedDateTime", "")[:10],
                body_text=body_text[:5000],
                folder_path=folder.full_path,
                has_attachments=msg.get("hasAttachments", False),
            ))
        url = data.get("@odata.nextLink", "")
    return messages


async def fetch_attachments(client: GraphClient, message: MessageData) -> None:
    data = await client.get(
        f"/users/{MAILBOX}/messages/{message.id}/attachments",
        params={"$select": "id,name,contentType,size"},
    )
    for att in data.get("value", []):
        name = att.get("name", "unknown")
        att_id = att.get("id", "")
        if not att_id or name.endswith(".ics"):
            continue

        try:
            att_full = await client.get(
                f"/users/{MAILBOX}/messages/{message.id}/attachments/{att_id}",
                params={"$select": "contentBytes"},
            )
            content_b64 = att_full.get("contentBytes", "")
        except Exception as e:
            logger.warning("Error downloading attachment %s: %s", name, e)
            continue

        if not content_b64:
            continue

        try:
            raw_bytes = base64.b64decode(content_b64)
            ext = Path(name).suffix.lower()
            text_content = None

            if ext == ".pdf":
                text_content = _extract_pdf(raw_bytes)
            elif ext in (".txt", ".csv", ".html", ".htm"):
                text_content = raw_bytes.decode("utf-8", errors="replace")[:5000]
            elif ext in (".xlsx", ".xls"):
                text_content = _extract_excel(raw_bytes)
            elif ext in (".doc", ".docx"):
                text_content = _extract_docx(raw_bytes)

            if text_content:
                message.attachments.append({
                    "name": name,
                    "text_content": text_content,
                })
        except Exception as e:
            logger.warning("Error decoding attachment %s: %s", name, e)


def _extract_pdf(raw_bytes: bytes) -> str | None:
    try:
        import pdfplumber, io
        with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
            texts = [p.extract_text() for p in pdf.pages[:20] if p.extract_text()]
            return "\n".join(texts)[:8000] if texts else None
    except ImportError:
        pass
    try:
        from PyPDF2 import PdfReader
        import io
        reader = PdfReader(io.BytesIO(raw_bytes))
        texts = [p.extract_text() for p in reader.pages[:20] if p.extract_text()]
        return "\n".join(texts)[:8000] if texts else None
    except ImportError:
        return None


def _extract_excel(raw_bytes: bytes) -> str | None:
    try:
        import openpyxl, io
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
        texts = []
        for ws in wb.worksheets[:5]:
            for row in ws.iter_rows(max_row=100, values_only=True):
                vals = [str(c) for c in row if c is not None]
                if vals:
                    texts.append(" | ".join(vals))
        return "\n".join(texts)[:5000] if texts else None
    except Exception:
        return None


def _extract_docx(raw_bytes: bytes) -> str | None:
    try:
        import docx, io
        doc = docx.Document(io.BytesIO(raw_bytes))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(texts)[:5000] if texts else None
    except Exception:
        return None


# ── Build email text for parsing ─────────────────────────────────


def build_email_text(msg: MessageData) -> str:
    """Build full text from email body + attachment contents."""
    parts = [
        f"From: {msg.sender_name} <{msg.sender_email}>",
        f"Date: {msg.date}",
        f"Subject: {msg.subject}",
        f"Folder: {msg.folder_path}",
        f"\n--- EMAIL BODY ---\n{msg.body_text}",
    ]
    for att in msg.attachments:
        if att.get("text_content"):
            parts.append(f"\n--- ATTACHMENT: {att['name']} ---\n{att['text_content']}")
    return "\n".join(parts)


# ── Main ─────────────────────────────────────────────────────────


async def main():
    graph = GraphClient()
    http = httpx.AsyncClient(base_url=API_BASE, timeout=120.0)

    try:
        # 1. Check Flask app is running
        try:
            r = await http.get("/partivia")
            r.raise_for_status()
        except Exception:
            print("ERROR: Flask app not running at", API_BASE)
            print("Start it first: python app.py")
            return

        # 2. Find PARTIVIA SPAGNA folder
        print("1/5  Finding PARTIVIA SPAGNA folder...")
        root = await find_partivia_folder(graph)
        if not root:
            print("ERROR: Folder not found!")
            return
        print(f"     Found: {root.full_path} ({root.message_count} messages)")

        # 3. Scan subfolders
        print("2/5  Scanning subfolders...")
        folders = await scan_subfolders(graph, root)
        total = sum(f.message_count for f in folders)
        print(f"     {len(folders)} folders, {total} messages total")
        for f in folders:
            print(f"     - {f.full_path} ({f.message_count} msg)")

        # 4. Download messages + attachments
        print("3/5  Downloading messages and attachments...")
        all_messages: list[MessageData] = []
        for folder in folders:
            msgs = await fetch_messages(graph, folder)
            for msg in msgs:
                if msg.has_attachments:
                    await fetch_attachments(graph, msg)
                    att_names = [a["name"] for a in msg.attachments]
                    if att_names:
                        print(f"     📎 {msg.subject[:60]} → {att_names}")
                    else:
                        print(f"     ✉  {msg.subject[:60]}")
                else:
                    print(f"     ✉  {msg.subject[:60]}")
            all_messages.extend(msgs)

        print(f"     Total: {len(all_messages)} messages, "
              f"{sum(len(m.attachments) for m in all_messages)} attachments with text")

        # 5. Confirm and clear existing data
        resp = input(f"\nReady to re-import {len(all_messages)} emails. "
                     f"This will DELETE all existing Partivia quotes. Continue? (y/N) ").strip().lower()
        if resp != 'y':
            print("Cancelled.")
            return

        print("4/5  Clearing existing quotes and email logs...")
        r = await http.delete("/api/partivia/quotes-all")
        if r.status_code == 200:
            print("     Quotes cleared.")
        else:
            print(f"     Warning: clear quotes returned {r.status_code}")

        # 6. Parse and import each email
        print("5/5  Parsing and importing emails...")
        stats = {"parsed": 0, "quotes": 0, "skipped": 0, "errors": 0}

        for i, msg in enumerate(all_messages):
            email_text = build_email_text(msg)
            label = f"[{i+1}/{len(all_messages)}] {msg.subject[:50]}"

            # Parse email
            try:
                r = await http.post("/api/partivia/parse-email",
                                    json={"text": email_text})
                j = r.json()
            except Exception as e:
                print(f"  ✗ {label} — parse error: {e}")
                stats["errors"] += 1
                continue

            if not j.get("ok"):
                print(f"  ✗ {label} — {j.get('error', 'unknown error')}")
                stats["errors"] += 1
                continue

            parsed = j.get("parsed", {})
            email_log_id = j.get("email_log_id")
            quotes = parsed.get("quotes", [])
            is_quote = parsed.get("is_quote", False)

            if not is_quote or not quotes:
                msg_type = parsed.get("message_type", "not a quote")
                print(f"  ○ {label} — {msg_type}")
                stats["skipped"] += 1
                stats["parsed"] += 1
                continue

            # Apply quotes
            try:
                r = await http.post("/api/partivia/apply",
                                    json={"quotes": quotes, "email_log_id": email_log_id})
                j2 = r.json()
            except Exception as e:
                print(f"  ✗ {label} — apply error: {e}")
                stats["errors"] += 1
                continue

            if j2.get("ok"):
                results = j2.get("results", [])
                added = sum(1 for r in results if r.get("action") == "added")
                updated = sum(1 for r in results if r.get("action") == "updated")
                hotels = [r.get("hotel", "?") for r in results]
                print(f"  ✓ {label} — {added} added, {updated} updated: {', '.join(hotels)}")
                stats["quotes"] += added + updated
            else:
                print(f"  ✗ {label} — apply failed: {j2.get('error')}")
                stats["errors"] += 1

            stats["parsed"] += 1

            # Small delay to avoid overwhelming the API
            await asyncio.sleep(1)

        # Report
        print(f"\n{'='*60}")
        print(f"DONE!")
        print(f"  Emails processed: {stats['parsed']}")
        print(f"  Quotes created:   {stats['quotes']}")
        print(f"  Skipped (not quotes): {stats['skipped']}")
        print(f"  Errors: {stats['errors']}")
        print(f"{'='*60}")

    finally:
        await graph.close()
        await http.aclose()


if __name__ == "__main__":
    asyncio.run(main())
